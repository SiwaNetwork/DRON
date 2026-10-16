"""Документы ЕСКД в Word (.docx, python-docx) — редактируемая копия PDF из eskd_pdf.py.

Источник тот же (Markdown docs/*.md, строки перечней из модели схемы), оформление то же:
* рамка листа (20/5 мм) и дополнительные графы 19–25 — в верхнем колонтитуле (рамка — фигура
  за текстом, графы — плавающие таблицы с текстом снизу вверх);
* основная надпись по ГОСТ 2.104: форма 2 (первый лист) и 2а (последующие) — таблицы в нижнем
  колонтитуле, номер листа и число листов — поля PAGE и NUMPAGES (Word считает сам);
* шрифт основных надписей osifont (tools/fonts) встроен в файл: Word на чужом ПК покажет его без
  установки; основной текст — Times New Roman, как в PDF.
Подлинник — PDF; правки вносятся в .md, DOCX пересобирается (tools/build_docs.py,
hw/gen/eskd_lists.py). Выход детерминирован: одинаковый исход -> побайтно одинаковый файл.
"""
import datetime
import io
import math
import os
import re
import uuid
import zipfile
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from reportlab.pdfbase import pdfmetrics

import eskd_pdf as E

NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
      'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" '
      'xmlns:w10="urn:schemas-microsoft-com:office:word"')
GOST = "osifont"                  # имя семейства во встроенном шрифте
SERIF, MONO = "Times New Roman", "Courier New"
THICK, THIN = 12, 6               # толщина линий, 1/8 пт: 1,5 пт ≈ 0,53 мм и 0,75 пт ≈ 0,26 мм
LINE_EM = (2697 + 709) / 2048     # высота строки osifont в Word (winAscent + winDescent), em
DATE = datetime.datetime(2026, 10, 1)


def tw(mm):
    """мм -> twips (1/20 пт)."""
    return int(round(mm * 1440 / 25.4))


def _pt(s, size_mm, avail_mm, row_mm=None):
    """Кегль, пт: высота шрифта size_mm (ГОСТ 2.304), уменьшенный по ширине строки s, как _fit в PDF;
    row_mm — высота строки таблицы с фиксированной высотой (строка Word должна в неё уместиться)."""
    pt = size_mm / 0.3528
    w1 = pdfmetrics.stringWidth(s, "GOST", 1.0) if s else 0
    if w1:
        pt = min(pt, avail_mm * 72 / 25.4 / w1)
    if row_mm:
        pt = min(pt, (row_mm * 72 / 25.4 - 0.5) / LINE_EM)
    return max(1.0, math.floor(pt * 2) / 2)


# ------------------------------------------------------------------ разметка WordprocessingML
def _rpr(font=None, pt=None, b=False, i=False, u=False, sup=False):
    x = ""
    if font:
        x += f'<w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:cs="{font}" w:eastAsia="{font}"/>'
    if b:
        x += "<w:b/><w:bCs/>"
    if i:
        x += "<w:i/><w:iCs/>"
    if pt:
        x += f'<w:sz w:val="{int(pt * 2)}"/><w:szCs w:val="{int(pt * 2)}"/>'
    if u:
        x += '<w:u w:val="single"/>'
    if sup:
        x += '<w:vertAlign w:val="superscript"/>'
    return f"<w:rPr>{x}</w:rPr>" if x else ""


def _r(text, rpr=""):
    """Прогон; \\n -> разрыв строки, \\t -> табуляция."""
    out = ""
    for k, line in enumerate(text.split("\n")):
        if k:
            out += "<w:br/>"
        for j, seg in enumerate(line.split("\t")):
            if j:
                out += "<w:tab/>"
            if seg:
                out += f'<w:t xml:space="preserve">{escape(seg)}</w:t>'
    return f"<w:r>{rpr}{out}</w:r>"


def _ppr(style=None, jc=None, before=0, after=0, line=240, rule="auto", ind=None, keep=False, tabs=None,
         mark=""):
    x = f'<w:pStyle w:val="{style}"/>' if style else ""
    if keep:
        x += "<w:keepNext/>"
    if tabs:
        x += "<w:tabs>" + "".join(f'<w:tab w:val="left" w:pos="{tw(t)}"/>' for t in tabs) + "</w:tabs>"
    if before is not None:
        x += f'<w:spacing w:before="{before}" w:after="{after}" w:line="{line}" w:lineRule="{rule}"/>'
    if ind is not None:
        x += ind
    if jc:
        x += f'<w:jc w:val="{jc}"/>'
    return f"<w:pPr>{x}{mark}</w:pPr>"


def _p(runs="", **kw):
    return f"<w:p>{_ppr(**kw)}{runs}</w:p>"


def _tiny():
    """Пустой абзац высотой 1 пт (разделитель плавающих таблиц, хвост колонтитула)."""
    return _p(line=20, rule="exact", ind='<w:ind w:left="0" w:right="0" w:firstLine="0"/>',
              mark=_rpr(pt=1))


NOIND = '<w:ind w:left="0" w:right="0" w:firstLine="0"/>'


def _cell_p(text, size_mm, w_mm, jc="center", row_mm=None, u=False, field=None, mar=0.5, after=0):
    """Абзац ячейки шрифтом osifont, кегль подобран по ширине ячейки w_mm за вычетом полей mar."""
    pt = _pt(text if not field else "99", size_mm, w_mm - 2 * mar, row_mm)
    rp = _rpr(GOST, pt, u=u)
    if field:   # составное поле: результат берёт шрифт прогона с кодом поля (у fldSimple — нет)
        runs = (f'<w:r>{rp}<w:fldChar w:fldCharType="begin"/></w:r>'
                f'<w:r>{rp}<w:instrText xml:space="preserve"> {field} </w:instrText></w:r>'
                f'<w:r>{rp}<w:fldChar w:fldCharType="separate"/></w:r>{_r("1", rp)}'
                f'<w:r>{rp}<w:fldChar w:fldCharType="end"/></w:r>')
    else:
        runs = _r(text, rp) if text else ""
    return _p(runs, jc=jc, after=after, ind=NOIND, mark=rp)


def _borders(top, left, bottom, right):
    def b(side, wgt):
        if not wgt:
            return f'<w:{side} w:val="nil"/>'
        return f'<w:{side} w:val="single" w:sz="{THICK if wgt == 2 else THIN}" w:space="0" w:color="000000"/>'
    return "<w:tcBorders>" + b("top", top) + b("left", left) + b("bottom", bottom) + b("right", right) + \
        "</w:tcBorders>"


def _grid(widths):
    """Ширины столбцов в twips с накоплением (сумма точно равна ширине таблицы)."""
    xs, acc = [0], 0.0
    for w in widths:
        acc += w
        xs.append(tw(acc))
    return [b - a for a, b in zip(xs, xs[1:])]


def grid_table(widths, heights, cells, hw, vw, ind=0, jc=None, pos=None, mar=0.5, hrule="exact",
               header_rows=0, head_mar=None, vtext=False):
    """Таблица ЕСКД с объединёнными ячейками.
    cells: {(строка, столбец): (строк, столбцов, xml абзацев)}; не покрытые ячейки — пустые 1×1.
    hw(b, c) — толщина горизонтальной линии над строкой b в столбце c (0 — нет, 1 — тонкая, 2 — основная),
    vw(k, r) — толщина вертикальной линии слева от столбца k в строке r.
    hrule — правило высоты строк ('exact', 'atLeast' или список по строкам); header_rows — повторяемая
    на каждом листе шапка (поля ячеек head_mar, мм); pos=(x, y) мм — плавающая таблица от края листа;
    vtext — текст снизу вверх (графы 19–25)."""
    hrules = hrule if isinstance(hrule, list) else [hrule] * len(heights)
    gw = _grid(widths)
    nr, nc = len(heights), len(widths)
    owner = {}
    for (r0, c0), (rs, cs, _) in cells.items():
        for r in range(r0, r0 + rs):
            for c in range(c0, c0 + cs):
                owner[(r, c)] = (r0, c0)
    for r in range(nr):
        for c in range(nc):
            if (r, c) not in owner:
                owner[(r, c)] = (r, c)
                cells[(r, c)] = (1, 1, "")
    pr = ""
    if pos:
        pr += (f'<w:tblpPr w:leftFromText="0" w:rightFromText="0" w:topFromText="0" w:bottomFromText="0" '
               f'w:vertAnchor="page" w:horzAnchor="page" w:tblpX="{tw(pos[0])}" w:tblpY="{tw(pos[1])}"/>'
               '<w:tblOverlap w:val="overlap"/>')
    pr += f'<w:tblW w:w="{sum(gw)}" w:type="dxa"/>'
    if jc:
        pr += f'<w:jc w:val="{jc}"/>'
    pr += f'<w:tblInd w:w="{tw(ind)}" w:type="dxa"/>'
    pr += '<w:tblLayout w:type="fixed"/>'
    pr += (f'<w:tblCellMar><w:top w:w="0" w:type="dxa"/><w:left w:w="{tw(mar)}" w:type="dxa"/>'
           f'<w:bottom w:w="0" w:type="dxa"/><w:right w:w="{tw(mar)}" w:type="dxa"/></w:tblCellMar>')
    pr += '<w:tblLook w:val="0000" w:firstRow="0" w:lastRow="0" w:firstColumn="0" w:lastColumn="0" ' \
          'w:noHBand="1" w:noVBand="1"/>'
    x = f"<w:tbl {NS}><w:tblPr>{pr}</w:tblPr><w:tblGrid>" + \
        "".join(f'<w:gridCol w:w="{g}"/>' for g in gw) + "</w:tblGrid>"
    for r in range(nr):
        x += (f'<w:tr><w:trPr><w:cantSplit/><w:trHeight w:val="{tw(heights[r])}" w:hRule="{hrules[r]}"/>'
              + ("<w:tblHeader/>" if r < header_rows else "") + "</w:trPr>")
        c = 0
        while c < nc:
            r0, c0 = owner[(r, c)]
            rs, cs, body = cells[(r0, c0)]
            cs_here = cs - (c - c0)
            last = r == r0 + rs - 1
            tc = f'<w:tcW w:w="{sum(gw[c:c + cs_here])}" w:type="dxa"/>'
            if cs_here > 1:
                tc += f'<w:gridSpan w:val="{cs_here}"/>'
            if rs > 1:
                tc += '<w:vMerge w:val="restart"/>' if r == r0 else "<w:vMerge/>"
            tc += _borders(hw(r, c) if r == r0 else 0, vw(c, r), hw(r + 1, c) if last else 0, vw(c + cs_here, r))
            if head_mar is not None and r < header_rows:
                tc += (f'<w:tcMar><w:left w:w="{tw(head_mar)}" w:type="dxa"/>'
                       f'<w:right w:w="{tw(head_mar)}" w:type="dxa"/></w:tcMar>')
            if vtext:
                tc += '<w:textDirection w:val="btLr"/>'
            tc += '<w:vAlign w:val="center"/>'
            x += f"<w:tc><w:tcPr>{tc}</w:tcPr>{body if r == r0 and body else _tiny()}</w:tc>"
            c += cs_here
        x += "</w:tr>"
    return x + "</w:tbl>"


# ------------------------------------------------------------------ рамка и основные надписи
def _vml_style(x, y, w, h, z):
    k = 72 / 25.4
    return (f"position:absolute;margin-left:{x * k:.2f}pt;margin-top:{y * k:.2f}pt;"
            f"width:{w * k:.2f}pt;height:{h * k:.2f}pt;z-index:{z};"
            "mso-position-horizontal:absolute;mso-position-horizontal-relative:page;"
            "mso-position-vertical:absolute;mso-position-vertical-relative:page")


def _frame(pw, ph, spid):
    """Рамка 20/5 мм: прямоугольник VML за текстом, привязан к листу."""
    return (f'<w:r><w:pict><v:rect id="eskd_frame{spid}" o:spid="_x0000_s{spid}" o:allowincell="f" filled="f" '
            f'strokeweight="1.5pt" style="{_vml_style(20, 5, pw - 25, ph - 10, -251659264)}">'
            '<w10:wrap anchorx="page" anchory="page"/></v:rect></w:pict></w:r>')


def _side(labels, heights, x, y, spid):
    """Дополнительные графы: столбцы 5 мм (название) и 7 мм, текст снизу вверх. Таблица — в надписи
    VML за текстом: плавающая таблица в колонтитуле вытесняла бы таблицу перечня, вплотную
    примыкающую к рамке. Надпись на 0,5 мм шире таблицы с каждой стороны, чтобы линии не срезались."""
    cells = {(i, 0): (1, 1, _cell_p(lab, 2.5, h)) for i, (lab, h) in enumerate(zip(labels, heights))}
    tbl = grid_table([5, 7], heights, cells, lambda b, c: 2, lambda k, r: 2, ind=0.5, mar=0, vtext=True)
    pad = _p(line=tw(0.5), rule="exact", ind=NOIND, mark=_rpr(pt=1))
    st = _vml_style(x - 0.5, y - 0.5, 13, sum(heights) + 1.5, -251658240 + spid)
    return (f'<w:r><w:pict><v:rect id="eskd_side{spid}" o:spid="_x0000_s{spid}" o:allowincell="f" filled="f" '
            f'stroked="f" style="{st}"><v:textbox inset="0,0,0,0"><w:txbxContent>{pad}{tbl}{_tiny()}'
            '</w:txbxContent></v:textbox><w10:wrap anchorx="page" anchory="page"/></v:rect></w:pict></w:r>')


def _header(pw, ph, first):
    s = 1025 if first else 1029
    # графы 19–23: «Подп. и дата», «Инв. № дубл.», «Взам. инв. №», «Подп. и дата», «Инв. № подл.»
    runs = _frame(pw, ph, s) + _side(["Подп. и дата", "Инв. № дубл.", "Взам. инв. №", "Подп. и дата",
                                      "Инв. № подл."], [35, 25, 25, 35, 25], 8, ph - 5 - 145, s + 1)
    if first:   # графы 24–25 первого листа
        runs += _side(["Перв. примен.", "Справ. №"], [60, 60], 8, 5, s + 2)
    return _p(runs, line=20, rule="exact", ind=NOIND, mark=_rpr(pt=1))


LEFT_W = [7, 10, 23, 15, 10]
LABELS = ["Изм.", "Лист", "№ докум.", "Подп.", "Дата"]


def stamp2(meta, **kw):
    """Форма 2 (185×40): 2 строки изменений, заголовки, 5 строк подписей; справа — обозначение,
    наименование, литера, лист/листов, предприятие. kw — положение таблицы (ind, jc)."""
    widths = LEFT_W + [70, 5, 5, 5, 15, 20]
    cells = {}
    for c, (w, lab) in enumerate(zip(LEFT_W, LABELS)):
        cells[(2, c)] = (1, 1, _cell_p(lab, 2.5, w, row_mm=5))
    vals = {"Разраб.": meta.get("dev", ""), "Пров.": meta.get("chk", ""), "Н.контр.": meta.get("nch", ""),
            "Утв.": meta.get("app", "")}
    for i, role in enumerate(["Разраб.", "Пров.", "", "Н.контр.", "Утв."]):
        r = 3 + i
        cells[(r, 0)] = (1, 2, _cell_p(role, 2.5, 17, jc="left", row_mm=5))
        cells[(r, 2)] = (1, 1, _cell_p(vals.get(role, ""), 2.5, 23, row_mm=5))
        if role == "Разраб." and meta.get("date"):
            d = meta["date"]
            cells[(r, 4)] = (1, 1, _cell_p(d[-5:] if len(d) > 8 else d, 2.5, 10, row_mm=5))
    cells[(0, 5)] = (3, 6, _cell_p(meta["designation"], 5, 120))
    cells[(3, 5)] = (5, 1, _cell_p(meta["title"], 3.5, 70, after=tw(3) if meta.get("doctype") else 0) +
                     (_cell_p(meta.get("doctype", ""), 3.5, 70) if meta.get("doctype") else ""))
    cells[(3, 6)] = (1, 3, _cell_p("Лит.", 2.5, 15, row_mm=5))
    cells[(3, 9)] = (1, 1, _cell_p("Лист", 2.5, 15, row_mm=5))
    cells[(3, 10)] = (1, 1, _cell_p("Листов", 2.5, 20, row_mm=5))
    cells[(4, 6)] = (1, 1, _cell_p(meta.get("lit", ""), 3, 5, row_mm=5))
    cells[(4, 9)] = (1, 1, _cell_p("", 3, 15, row_mm=5, field="PAGE"))
    cells[(4, 10)] = (1, 1, _cell_p("", 3, 20, row_mm=5, field="NUMPAGES"))
    cells[(5, 6)] = (3, 5, _cell_p(meta.get("company", ""), 3.5, 50))

    def hw(b, c):     # левая часть: тонкие — между строками изменений и между строками подписей
        return 1 if c <= 4 and b in (1, 4, 5, 6, 7) else 2

    def vw(k, r):     # тонкие — деления графы «Лит.»
        return 1 if k in (7, 8) else 2
    return grid_table(widths, [5] * 8, cells, hw, vw, mar=0.5, **kw)


def stamp2a(meta, **kw):
    """Форма 2а (185×15): строки 5 + (2 + 3) + 5 мм — деление графы «Лист» на 7/8 мм."""
    widths = LEFT_W + [110, 10]
    cells = {}
    for c, (w, lab) in enumerate(zip(LEFT_W, LABELS)):
        cells[(1, c)] = (2, 1, "")
        cells[(3, c)] = (1, 1, _cell_p(lab, 2.5, w, row_mm=5))
    cells[(0, 5)] = (4, 1, _cell_p(meta["designation"], 5, 110))
    cells[(0, 6)] = (2, 1, _cell_p("Лист", 2.5, 10))
    cells[(2, 6)] = (2, 1, _cell_p("", 3.5, 10, field="PAGE"))
    return grid_table(widths, [5, 2, 3, 5], cells, lambda b, c: 1 if c <= 4 and b == 1 else 2,
                      lambda k, r: 2, mar=0.5, **kw)


def _footer(meta, pw, left, first, spacer, stamp_jc=None, ind=0):
    """Основная надпись + «Копировал» / «Формат» под рамкой. Колонтитул кончается на 1,5 мм от края
    листа: строка «Копировал» 3,5 мм, низ основной надписи — на линии рамки (5 мм)."""
    R = pw - 5
    x = _p(line=tw(5), rule="exact", ind=NOIND, mark=_rpr(pt=1)) if spacer else ""
    x += (stamp2 if first else stamp2a)(meta, ind=ind, jc=stamp_jc)
    rp = _rpr(GOST, 7)
    x += _p(_r(f"\tКопировал\tФормат {meta.get('format', 'A4')}", rp), line=tw(3.5), rule="exact",
            ind=NOIND, tabs=[R - 110 - left, R - 40 - left], mark=rp)
    return x


def _set_hf(part, xml):
    el = part._element
    for ch in list(el):
        el.remove(ch)
    for ch in parse_xml(f"<w:hdr {NS}>{xml}</w:hdr>"):
        el.append(ch)


# ------------------------------------------------------------------ документ
def _put(parent, child, order):
    """Вставить child в parent с учётом порядка элементов схемы (order — локальные имена)."""
    name = child.tag.split("}")[1]
    for old in parent.findall(qn(f"w:{name}")):
        parent.remove(old)
    after = order[order.index(name) + 1:]
    for ch in parent:
        if ch.tag.split("}")[1] in after:
            ch.addprevious(child)
            return child
    parent.append(child)
    return child


RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
             "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
             "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
             "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
             "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr"]


def _w(xml):
    """Элемент из фрагмента разметки: объявления пространств имён — в первый тег."""
    if " xmlns:w=" in xml.split(">", 1)[0]:
        return parse_xml(xml)
    tag = re.match(r"<[\w:]+", xml).group(0)
    return parse_xml(xml.replace(tag, f"{tag} {NS}", 1))


def _style_props(st, ppr=(), rpr=()):
    el = st.element
    p = el.find(qn("w:pPr"))
    if p is None:
        p = _put(el, parse_xml(f"<w:pPr {NS}/>"), ["name", "aliases", "basedOn", "next", "link", "autoRedefine",
                                                  "hidden", "uiPriority", "semiHidden", "unhideWhenUsed",
                                                  "qFormat", "locked", "personal", "personalCompose",
                                                  "personalReply", "rsid", "pPr", "rPr", "tblPr", "trPr", "tcPr",
                                                  "tblStylePr"])
    r = el.find(qn("w:rPr"))
    if r is None:
        r = parse_xml(f"<w:rPr {NS}/>")
        p.addnext(r)
    for ch in list(r):
        if ch.tag in (qn("w:rFonts"), qn("w:color")):
            r.remove(ch)
    for x in ppr:
        _put(p, _w(x), PPR_ORDER)
    for x in rpr:
        _put(r, _w(x), RPR_ORDER)


def _new_doc():
    doc = Document()
    st = doc.styles.element
    dd = st.find(qn("w:docDefaults"))
    rd = dd.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    _put(rd, _w(f'<w:rFonts w:ascii="{SERIF}" w:hAnsi="{SERIF}" w:cs="{SERIF}" w:eastAsia="{SERIF}"/>'), RPR_ORDER)
    _put(rd, _w('<w:lang w:val="ru-RU" w:eastAsia="ru-RU" w:bidi="ar-SA"/>'), RPR_ORDER)
    pd = dd.find(qn("w:pPrDefault")).find(qn("w:pPr"))
    _put(pd, _w('<w:spacing w:after="0" w:line="240" w:lineRule="auto"/>'), PPR_ORDER)
    # основной текст (ГОСТ 2.105): абзацный отступ 12,5 мм, по ширине, интервал как в PDF (1,35 кегля)
    body = doc.styles.add_style("Текст ЕСКД", WD_STYLE_TYPE.PARAGRAPH)
    body.style_id = "EskdText"
    body.base_style = doc.styles["Normal"]
    body.quick_style = True
    _style_props(body, [f'<w:spacing w:after="40" w:line="282" w:lineRule="auto"/>',
                        f'<w:ind w:firstLine="{tw(12.5)}"/>', '<w:jc w:val="both"/>'], ['<w:sz w:val="22"/>'])
    for lvl, (pt, bef, aft) in ((1, (14, 8, 6)), (2, (12, 6, 4)), (3, (11, 4, 3))):
        _style_props(doc.styles[f"Heading {lvl}"],
                     [f'<w:spacing w:before="{bef * 20}" w:after="{aft * 20}" w:line="271" w:lineRule="auto"/>',
                      f'<w:ind w:firstLine="{tw(12.5)}"/>'],
                     [f'<w:rFonts w:ascii="{SERIF}" w:hAnsi="{SERIF}" w:cs="{SERIF}" w:eastAsia="{SERIF}"/>',
                      '<w:color w:val="000000"/>', f'<w:sz w:val="{pt * 2}"/>', f'<w:szCs w:val="{pt * 2}"/>'])
    # совместимость Word 2013+ (отступ таблиц от края рамки, а не от текста ячейки); встроенные шрифты
    s = doc.settings.element
    for cs in s.iter(qn("w:compatSetting")):
        if cs.get(qn("w:name")) == "compatibilityMode":
            cs.set(qn("w:val"), "15")
    s.find(qn("w:zoom")).addnext(_w("<w:embedTrueTypeFonts/>"))
    _embed_font(doc)
    pkg = doc.part.package
    for rid, rel in list(pkg.rels.items()):      # эскиз пустой страницы шаблона python-docx
        if rel.reltype == RT.THUMBNAIL:
            pkg.rels.pop(rid)
    return doc


def _embed_font(doc):
    """osifont во встроенном виде (ECMA-376, 17.8.1): первые 32 байта шрифта XOR ключ из GUID."""
    ft = doc.part.part_related_by(RT.FONT_TABLE)
    guid = str(uuid.uuid5(uuid.NAMESPACE_URL, "https://github.com/SiwaNetwork/DRON/osifont-mgp8")).upper()
    key = bytes.fromhex(guid.replace("-", ""))[::-1]
    data = bytearray(open(E.FONT_FILE, "rb").read())
    for i in range(32):
        data[i] ^= key[i % 16]
    part = Part(PackURI("/word/fonts/font1.odttf"), "application/vnd.openxmlformats-officedocument.obfuscatedFont",
                bytes(data), doc.part.package)
    rid = ft.relate_to(part, RT.FONT)
    entry = (f'<w:font w:name="{GOST}"><w:charset w:val="00"/><w:family w:val="swiss"/>'
             f'<w:pitch w:val="variable"/><w:embedRegular r:id="{rid}" w:fontKey="{{{guid}}}"/></w:font>')
    ft._blob = ft.blob.replace(b"</w:fonts>", entry.encode("utf-8") + b"</w:fonts>")


def _section(doc, pw, ph, top, bottom, left, right, footer_mm):
    sec = doc.sections[0]
    if pw > ph:
        sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(pw), Mm(ph)
    sec.top_margin, sec.bottom_margin = Mm(top), Mm(bottom)
    sec.left_margin, sec.right_margin = Mm(left), Mm(right)
    sec.header_distance, sec.footer_distance = Mm(0), Mm(footer_mm)
    sec.gutter = Mm(0)
    sec.different_first_page_header_footer = True
    return sec


def _furniture(doc, sec, meta, pw, ph, left, spacer, stamp_jc=None, ind=0):
    for hdr, first in ((sec.first_page_header, True), (sec.header, False)):
        hdr.is_linked_to_previous = False
        _set_hf(hdr, _header(pw, ph, first))
    for ftr, first in ((sec.first_page_footer, True), (sec.footer, False)):
        ftr.is_linked_to_previous = False
        _set_hf(ftr, _footer(meta, pw, left, first, spacer, stamp_jc, ind))


def _save(doc, path, meta):
    cp = doc.core_properties
    cp.title = meta.get("doctype") or meta.get("title", "")
    cp.subject = meta.get("designation", "")
    cp.author = cp.last_modified_by = meta.get("company", "")
    cp.comments = ""
    cp.revision = 1
    cp.created = cp.modified = DATE
    buf = io.BytesIO()
    doc.save(buf)
    # пересборка архива с постоянными датами: одинаковый исход -> одинаковый файл
    zin = zipfile.ZipFile(buf)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            zout.writestr(zi, zin.read(info.filename))


# ------------------------------------------------------------------ текстовые документы
def _inline(t, base=None):
    """Строчная разметка как в eskd_pdf.inline(): `код`, **полужирный**, *курсив*, [текст](адрес), 2^32."""
    base = dict(base or {})
    out = ""
    for k, part in enumerate(re.split(r"`([^`]+)`", t)):
        if k % 2:
            out += _r(part, _rpr(MONO, base.get("pt"), base.get("b"), base.get("i")))
            continue
        part = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", part)
        pos = 0
        for m in re.finditer(r"\*\*(.+?)\*\*|(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?!\w)", part):
            out += _sup(part[pos:m.start()], base)
            if m.group(1) is not None:
                out += _inline(m.group(1), dict(base, b=True))
            else:
                out += _inline(m.group(2), dict(base, i=True))
            pos = m.end()
        out += _sup(part[pos:], base)
    return out


def _sup(t, base):
    out, pos = "", 0
    for m in re.finditer(r"(\d|\))\^([−-]?[0-9A-Za-zμα-ω]+)", t):
        out += _r(t[pos:m.end(1)], _rpr(None, base.get("pt"), base.get("b"), base.get("i")))
        out += _r(m.group(2), _rpr(None, base.get("pt"), base.get("b"), base.get("i"), sup=True))
        pos = m.end()
    if t[pos:]:
        out += _r(t[pos:], _rpr(None, base.get("pt"), base.get("b"), base.get("i")))
    return out


def _md_table(rows, width_mm):
    gw = _grid(E.table_widths(rows, width_mm))
    bd = "".join(f'<w:{s} w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                 for s in ("top", "left", "bottom", "right", "insideH", "insideV"))
    x = (f'<w:tbl {NS}><w:tblPr><w:tblW w:w="{sum(gw)}" w:type="dxa"/><w:tblInd w:w="0" w:type="dxa"/>'
         f'<w:tblBorders>{bd}</w:tblBorders><w:tblLayout w:type="fixed"/>'
         f'<w:tblCellMar><w:top w:w="{tw(0.5)}" w:type="dxa"/><w:left w:w="{tw(1.5)}" w:type="dxa"/>'
         f'<w:bottom w:w="{tw(0.5)}" w:type="dxa"/><w:right w:w="{tw(1.5)}" w:type="dxa"/></w:tblCellMar>'
         '<w:tblLook w:val="0000" w:firstRow="0" w:lastRow="0" w:firstColumn="0" w:lastColumn="0" '
         'w:noHBand="1" w:noVBand="1"/></w:tblPr><w:tblGrid>'
         + "".join(f'<w:gridCol w:w="{g}"/>' for g in gw) + "</w:tblGrid>")
    for ri, r in enumerate(rows):
        x += "<w:tr><w:trPr><w:cantSplit/>" + ("<w:tblHeader/>" if ri == 0 else "") + "</w:trPr>"
        for g, c in zip(gw, r):
            x += (f'<w:tc><w:tcPr><w:tcW w:w="{g}" w:type="dxa"/><w:vAlign w:val="center"/></w:tcPr>'
                  + _p(_inline(c, {"pt": 9}), jc="center" if ri == 0 else "left", ind=NOIND,
                       mark=_rpr(pt=9)) + "</w:tc>")
        x += "</w:tr>"
    return x + "</w:tbl>"


def md_to_docx(md_path, docx_path, meta, width_mm=175):
    """Текстовый документ (ПЗ, ПМ, ТЗ, описание программы): A4, поля как у PDF (текст 25…200 мм)."""
    md = open(md_path, encoding="utf-8").read()
    basedir = os.path.dirname(md_path)
    doc = _new_doc()
    sec = _section(doc, 210, 297, top=15, bottom=25, left=25, right=10, footer_mm=1.5)
    _furniture(doc, sec, meta, 210, 297, left=25, spacer=True, ind=-5)
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))

    def add(xml):
        sect.addprevious(_w(xml))
    for b in E.md_blocks(md):
        kind = b[0]
        if kind == "p":
            add(_p(_inline(b[1]), style="EskdText", before=None))
        elif kind == "h":
            add(_p(_inline(b[2]), style=f"Heading{b[1]}", before=None))
        elif kind == "li":
            _, bullet, ind, text = b
            add(_p(_r(bullet + "\t") + _inline(text), style="EskdText", before=None,
                   ind=f'<w:ind w:left="{tw(17 + 5 * ind)}" w:hanging="{tw(4.5)}"/>'))
        elif kind == "tcap":
            add(_p(_inline(b[1]), style="EskdText", before=80, after=40, line=282, ind=NOIND, jc="left",
                   keep=True))
        elif kind == "code":
            add(_p(_r(b[1], _rpr(MONO, 8)), after=60, ind=NOIND, jc="left", mark=_rpr(MONO, 8)))
        elif kind == "table":
            add(_md_table(b[1], width_mm))
            add(_p(line=80, rule="exact", ind=NOIND, mark=_rpr(pt=1)))
        elif kind == "img":
            path = os.path.join(basedir, b[2])
            if os.path.exists(path):
                w, h = E.image_size(path, width_mm)
                p = doc.add_paragraph()
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.first_line_indent = Mm(0)
                p.add_run().add_picture(path, width=Mm(w), height=Mm(h))
                add(_p(_inline(b[1]), style="EskdText", before=40, after=120, line=282, ind=NOIND, jc="center"))
            else:
                add(_p(_r(f"[нет файла {b[2]}]"), style="EskdText", before=None, ind=NOIND))
    _save(doc, docx_path, meta)


# ------------------------------------------------------------------ перечни
# Запас до основной надписи на каждом листе перечня, мм: Word 2016 переносит последнюю строку листа
# на следующий, если остаётся меньше ≈ 1,05 мм (опыт: 1,0 — перенос, 1,1 — нет)
LIST_GAP = 1.2


def list_to_docx(path, pw, ph, meta, rows, widths, header, aligns, row_h=8, head_h=15, font=3.5):
    """Перечень элементов, спецификация, ведомость: таблица на всю рамку с повторяющейся шапкой.
    Строк на листе столько же, как в PDF (list_pages): 29 на первом листе A4 и 32 на последующих;
    последний лист добит пустыми строками."""
    doc = _new_doc()
    sec = _section(doc, pw, ph, top=5, bottom=20, left=20, right=5, footer_mm=1.5)
    _furniture(doc, sec, meta, pw, ph, left=20, spacer=False, stamp_jc="right" if pw > 210 else None)
    first = int((ph - 10 - 40 - head_h) // row_h)
    nxt = int((ph - 10 - 15 - head_h) // row_h)
    pages = 1 if len(rows) <= first else 1 + math.ceil((len(rows) - first) / nxt)
    rows = list(rows) + [[""] * len(widths)] * (first + (pages - 1) * nxt - len(rows))
    cells = {}
    for k, h in enumerate(header):   # шапка: строки через 4 мм, кегль 3 мм, поля ячеек 0,5 мм
        cells[(0, k)] = (1, 1, "".join(_p(_r(ln, _rpr(GOST, _pt(ln, 3.0, widths[k] - 1))), jc="center",
                                          line=tw(4.0), rule="exact", ind=NOIND, mark=_rpr(GOST, 8))
                                       for ln in h.split("\n")))
    for j, r in enumerate(rows, start=1):
        for k, val in enumerate(r):
            if isinstance(val, tuple):     # (текст, 'u') — подчёркнутый заголовок группы
                cells[(j, k)] = (1, 1, _cell_p(val[0], font, widths[k], u=True, mar=1.2))
            elif val:
                cells[(j, k)] = (1, 1, _cell_p(val, font, widths[k], jc=aligns[k], mar=1.2))
    # строки — «не менее 8 мм» (в Word можно дописать текст), шапка — ровно 15 мм. К высоте строки
    # «не менее» Word прибавляет толщину её верхней линии — вычитаем её (и 0,01 мм на округление).
    # Последняя строка каждого листа добирает остаток высоты до основной надписи без LIST_GAP
    # (на листе 1 она на ≈ 0,9 мм короче остальных, на листах 2… — на 0,1 мм длиннее)
    line_mm = {2: THICK / 8 * 25.4 / 72, 1: THIN / 8 * 25.4 / 72}
    pitch = row_h - 0.01
    heights = [head_h] + [pitch - line_mm[2 if j == 1 else 1] for j in range(1, len(rows) + 1)]
    for p in range(pages):
        n, avail = (first, ph - 10 - 40 - head_h) if p == 0 else (nxt, ph - 10 - 15 - head_h)
        heights[first + p * nxt] += avail - n * pitch - LIST_GAP
    tbl = grid_table(widths, heights, cells,
                     lambda b, c: 2 if b <= 1 or b == len(rows) + 1 else 1, lambda k, r: 2, mar=1.2,
                     hrule=["exact"] + ["atLeast"] * len(rows), header_rows=1, head_mar=0.5)
    sect = doc.element.body.find(qn("w:sectPr"))
    sect.addprevious(parse_xml(tbl))
    sect.addprevious(parse_xml(f"<w:p {NS}>" + _ppr(line=2, rule="exact", ind=NOIND, mark=_rpr(pt=1)) + "</w:p>"))
    _save(doc, path, meta)
    return pages
